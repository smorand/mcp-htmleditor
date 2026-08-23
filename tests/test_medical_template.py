"""Tests for the 'medical' slide charter: template, layouts, PPTX export.

Three axes, matching the three ways the charter can silently break:

* the CSS is duplicated in the reference decks (they must open standalone), so
  it is checked byte for byte against its single source, the bootstrap;
* the insertable layouts live in JavaScript, out of reach of the Python tests,
  so their required attributes are checked textually;
* the exporter has a dedicated chrome and theme, so a real export of the
  reference decks is run and inspected.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from bs4 import BeautifulSoup, Tag
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from mcp_htmleditor.export.pptx_style import StyleResolver, detect_theme
from mcp_htmleditor.export.to_pptx import find_slides, to_pptx
from mcp_htmleditor.templates import template_path

REPO_ROOT = Path(__file__).resolve().parents[1]
REFERENCE_DIR = REPO_ROOT / "templates" / "reference" / "slides"
LAYOUTS_JS = REPO_ROOT / "src" / "mcp_htmleditor" / "static" / "slide-layouts.js"
CATALOGUE = REFERENCE_DIR / "medical.html"
EXAMPLE = REFERENCE_DIR / "example-medical-complete.html"

# Anchored at line start: a header comment mentioning the style tag would
# otherwise open the match and the captured "CSS" would include <html>/<head>.
_STYLE_RE = re.compile(r"^<style>\n.*?^</style>", re.DOTALL | re.MULTILINE)

MEDICAL_LAYOUTS = (
    "title",
    "disclosure",
    "agenda",
    "section",
    "content",
    "image",
    "grid",
    "compare",
    "columns",
    "case",
    "steps",
    "table",
    "keymessage",
    "takehome",
    "references",
    "thanks",
)


def _style_block(path: Path) -> str:
    """Return the first ``<style>`` block of an HTML file."""
    match = _STYLE_RE.search(path.read_text(encoding="utf-8"))
    assert match is not None, f"no <style> block in {path}"
    return match.group(0)


def _medical_layouts_js() -> str:
    """Return the ``LAYOUT_SETS.medical`` part of slide-layouts.js."""
    text = LAYOUTS_JS.read_text(encoding="utf-8")
    start = text.index("LAYOUT_SETS.medical")
    return text[start:]


# ---------------------------------------------------------------------------
# Bootstrap
# ---------------------------------------------------------------------------


def test_medical_bootstrap_is_a_presentation_with_the_charter_markers() -> None:
    """The 'medical' bootstrap declares the charter and its detection tokens."""
    content = template_path("medical").read_text(encoding="utf-8")

    assert 'data-doc-type="presentation"' in content
    assert 'data-doc-template="medical"' in content
    assert "--med-teal:#159984" in content
    assert "--med-orange:#EE5A02" in content


def test_medical_bootstrap_ships_fullscreen_support() -> None:
    """Every slide template must ship the presentation mode (see AGENTS.md)."""
    content = template_path("medical").read_text(encoding="utf-8")

    for marker in (":fullscreen", 'id="btn-present"', "enterPresentation", "exitPresentation", "updateFullscreenScale"):
        assert marker in content, f"missing fullscreen marker {marker!r}"


def test_medical_bootstrap_cover_carries_the_charter_signature() -> None:
    """The cover ships the grey band, the two segment rule and an author block."""
    soup = BeautifulSoup(template_path("medical").read_text(encoding="utf-8"), "html.parser")
    slides = find_slides(soup)

    assert len(slides) == 1
    cover = slides[0]
    assert cover.get("data-slide-type") == "title"
    assert cover.find(class_="med-band") is not None
    rule = cover.find(class_="med-rule")
    assert isinstance(rule, Tag)
    assert [child.name for child in rule.find_all(True)] == ["i", "b"]
    assert cover.find(class_="author-name") is not None


# ---------------------------------------------------------------------------
# Reference decks
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("deck", [CATALOGUE, EXAMPLE], ids=["catalogue", "example"])
def test_reference_decks_reuse_the_bootstrap_css_verbatim(deck: Path) -> None:
    """A reference deck embeds the bootstrap CSS unchanged (make sync-medical-css)."""
    assert deck.is_file(), f"missing reference deck {deck}"

    assert _style_block(deck) == _style_block(template_path("medical")), (
        f"{deck.name} CSS diverged from the bootstrap: run `make sync-medical-css`"
    )


@pytest.mark.parametrize("deck", [CATALOGUE, EXAMPLE], ids=["catalogue", "example"])
def test_reference_decks_are_internally_coherent(deck: Path) -> None:
    """Slide ids, TOTAL, slideNames and page numbers stay in sync in a deck."""
    text = deck.read_text(encoding="utf-8")
    soup = BeautifulSoup(text, "html.parser")
    slides = find_slides(soup)

    assert len(slides) >= 8
    for index, slide in enumerate(slides):
        assert slide.get("id") == f"slide-{index}"
        assert slide.get("data-id") == f"slide-{index}"
        assert str(slide.get("data-title") or "").strip(), f"slide-{index} has no data-title"
        assert str(slide.get("data-slide-type") or "").strip(), f"slide-{index} has no data-slide-type"

    total = re.search(r"const TOTAL = (\d+);", text)
    assert total is not None
    assert int(total.group(1)) == len(slides)

    names = re.search(r"const slideNames = \[(.*?)\];", text, re.DOTALL)
    assert names is not None
    listed = re.findall(r'"((?:[^"\\]|\\.)*)"', names.group(1))
    assert listed == [str(slide.get("data-title")) for slide in slides]


@pytest.mark.parametrize("deck", [CATALOGUE, EXAMPLE], ids=["catalogue", "example"])
def test_reference_decks_cite_every_image_bearing_slide(deck: Path) -> None:
    """A slide carrying an image carries a non empty source line."""
    soup = BeautifulSoup(deck.read_text(encoding="utf-8"), "html.parser")

    for slide in find_slides(soup):
        if not slide.find("img"):
            continue
        sources = [
            source.get_text(strip=True) for source in slide.find_all(class_="med-source") if source.get_text(strip=True)
        ]
        assert sources, f"slide {slide.get('id')} shows an image without a .med-source line"


@pytest.mark.parametrize("deck", [CATALOGUE, EXAMPLE], ids=["catalogue", "example"])
def test_reference_decks_never_use_a_css_grid_row(deck: Path) -> None:
    """Rows are flex with explicit w-* widths: a CSS grid exports as a stack."""
    text = deck.read_text(encoding="utf-8")

    assert "grid-template-columns" not in text
    assert 'class="med-grid cols-' not in text


# ---------------------------------------------------------------------------
# Insertable layouts (browser side)
# ---------------------------------------------------------------------------


def test_layouts_js_exposes_every_medical_layout() -> None:
    """The picker offers the whole medical catalogue of slide types."""
    block = _medical_layouts_js()

    for key in MEDICAL_LAYOUTS:
        assert re.search(rf"^  {key}: \{{", block, re.MULTILINE), f"missing layout {key!r}"


def test_layouts_js_medical_fragments_carry_the_required_attributes() -> None:
    """Each fragment is a slide article and keeps the renumbered page marker."""
    block = _medical_layouts_js()
    fragments = re.findall(r"html: `\s*(<article.*?</article>)`", block, re.DOTALL)

    assert len(fragments) == len(MEDICAL_LAYOUTS)
    for fragment in fragments:
        assert 'data-type="slide"' in fragment
        assert 'data-id="{{ID}}"' in fragment
        assert "data-slide-type=" in fragment
        assert "data-title=" in fragment
        # cover, section separator and closing slide carry no footer band.
        if any(f'data-slide-type="{kind}"' in fragment for kind in ("title", "section", "thanks")):
            continue
        assert 'class="slide-foot-page"' in fragment


def test_layouts_js_detects_the_medical_template_first() -> None:
    """detectTemplate must test the medical markers before the EI ones."""
    text = LAYOUTS_JS.read_text(encoding="utf-8")
    detect = text[text.index("function detectTemplate") : text.index("function getLayouts")]

    assert detect.index("--med-teal") < detect.index("--ei-blue")


# ---------------------------------------------------------------------------
# PPTX export
# ---------------------------------------------------------------------------


def test_detect_theme_returns_the_medical_charter() -> None:
    """A medical document resolves to the medical theme, not carbon nor generic."""
    soup = BeautifulSoup(template_path("medical").read_text(encoding="utf-8"), "html.parser")

    theme = detect_theme(soup)

    assert theme.key == "medical"
    assert theme.primary_alt == "159984"
    assert theme.accent == "EE5A02"


def test_medical_dark_slide_resolves_its_inverted_palette() -> None:
    """A `.slide.dark` caption resolves to the dark text colour, not the light one."""
    soup = BeautifulSoup(template_path("medical").read_text(encoding="utf-8"), "html.parser")
    resolver = StyleResolver.from_soup(soup)
    caption = BeautifulSoup('<figcaption class="med-caption">x</figcaption>', "html.parser").figcaption
    assert caption is not None

    assert resolver.color(caption, ("color",)) == "5D5D5D"
    resolver.active_scope = frozenset({"slide", "dark"})
    assert resolver.color(caption, ("color",)) == "9AA7B4"


@pytest.mark.parametrize("deck", [CATALOGUE, EXAMPLE], ids=["catalogue", "example"])
def test_medical_decks_export_every_slide_without_warning(deck: Path, tmp_path: Path) -> None:
    """The reference decks export one PPTX slide per HTML slide, cleanly."""
    expected = len(find_slides(BeautifulSoup(deck.read_text(encoding="utf-8"), "html.parser")))
    out = tmp_path / f"{deck.stem}.pptx"

    report = to_pptx(str(deck), str(out))

    assert report.charter == "medical"
    assert report.slide_count == expected
    assert report.warnings == []
    assert len(Presentation(str(out)).slides) == expected


def test_medical_export_keeps_pictures_tables_and_page_numbers(tmp_path: Path) -> None:
    """The exported deck keeps its images, its native tables and its page numbers."""
    out = tmp_path / "catalogue.pptx"
    to_pptx(str(CATALOGUE), str(out))
    prs = Presentation(str(out))

    pictures = sum(1 for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    tables = sum(1 for slide in prs.slides for shape in slide.shapes if shape.has_table)
    texts = [
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]

    assert pictures >= 6
    assert tables >= 2
    assert any(text.strip().isdigit() for text in texts), "no bare page number exported"
